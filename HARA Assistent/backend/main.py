"""
FastAPI backend — all routes for the HARA workflow.
"""

import os
import uuid
import shutil
import tempfile
from io import BytesIO
from typing import Dict, Any, List

from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from dotenv import load_dotenv

from .rag_engine import extract_pdf_content, build_faiss_store, load_faiss_store, retrieve_context, format_context
from .hara_chain import suggest_hazards, suggest_sec, draft_safety_goals
from .asil_rules import compute_asil, asil_rationale

load_dotenv()

app = FastAPI(title="HARA Assistant API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:8501", "http://127.0.0.1:8501"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# In-memory session store: session_id → faiss_path
sessions: Dict[str, str] = {}


# ── Pydantic Models ───────────────────────────────────────────────────────────

class HazardRequest(BaseModel):
    session_id: str
    system_description: str

class SECRequest(BaseModel):
    session_id: str
    hazardous_event: str
    operational_situation: str

class ASILRequest(BaseModel):
    S: str
    E: str
    C: str

class GoalRequest(BaseModel):
    session_id: str
    hazardous_event: str
    asil: str
    affected_function: str

class TraceabilityRow(BaseModel):
    id: str
    hazard: str
    hazardous_event: str
    operational_situation: str
    affected_function: str
    source_evidence: str
    S: str
    E: str
    C: str
    asil: str
    safety_goal: str
    fsr: str
    tsr: str
    fsr_status: str   # "Approved" | "Rejected" | "Edited"
    tsr_status: str

class ExportRequest(BaseModel):
    rows: List[TraceabilityRow]


# ── Routes ────────────────────────────────────────────────────────────────────

@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/ingest")
async def ingest_pdf(file: UploadFile = File(...)):
    """
    Accept a PDF upload, extract content, build FAISS index.
    Returns a session_id for subsequent requests.
    """
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Only PDF files are accepted.")

    session_id = str(uuid.uuid4())
    faiss_path = os.path.join(tempfile.gettempdir(), f"hara_{session_id}")

    # Save uploaded file to temp location
    tmp_pdf = os.path.join(tempfile.gettempdir(), f"{session_id}.pdf")
    try:
        content = await file.read()
        with open(tmp_pdf, "wb") as f:
            f.write(content)

        documents = extract_pdf_content(tmp_pdf)
        if not documents:
            raise HTTPException(status_code=422, detail="No extractable content found in PDF.")

        build_faiss_store(documents, faiss_path)
        sessions[session_id] = faiss_path

    finally:
        if os.path.exists(tmp_pdf):
            os.remove(tmp_pdf)

    return {
        "session_id": session_id,
        "pages_extracted": len(documents),
        "message": "PDF ingested successfully.",
    }


@app.post("/suggest-hazards")
def suggest_hazards_route(req: HazardRequest):
    """Retrieve context and suggest hazards using LLM."""
    if req.session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found. Please upload a PDF first.")

    store = load_faiss_store(sessions[req.session_id])
    docs = retrieve_context(store, req.system_description + " hazard failure mode risk", k=6)
    context = format_context(docs)

    hazards = suggest_hazards(context, req.system_description)
    return {"hazards": hazards, "context_used": context[:500] + "..."}


@app.post("/suggest-sec")
def suggest_sec_route(req: SECRequest):
    """Retrieve context and suggest S/E/C values using LLM."""
    if req.session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found.")

    store = load_faiss_store(sessions[req.session_id])
    query = f"{req.hazardous_event} {req.operational_situation} severity exposure controllability"
    docs = retrieve_context(store, query, k=4)
    context = format_context(docs)

    sec = suggest_sec(context, req.hazardous_event, req.operational_situation)
    return sec


@app.post("/compute-asil")
def compute_asil_route(req: ASILRequest):
    """Compute ASIL deterministically — no LLM."""
    try:
        level = compute_asil(req.S, req.E, req.C)
        rationale = asil_rationale(req.S, req.E, req.C)
        return {"asil": level, "asil_full": f"ASIL {level}" if level != "QM" else "QM", "rationale": rationale}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/draft-goals")
def draft_goals_route(req: GoalRequest):
    """Retrieve context and draft safety goals using LLM."""
    if req.session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found.")

    store = load_faiss_store(sessions[req.session_id])
    query = f"{req.hazardous_event} {req.affected_function} safety requirement"
    docs = retrieve_context(store, query, k=5)
    context = format_context(docs)

    goals = draft_safety_goals(context, req.hazardous_event, req.asil, req.affected_function)
    return goals


@app.post("/export-traceability")
def export_traceability(req: ExportRequest):
    """Build and return a styled Excel traceability matrix."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HARA Traceability Matrix"

    # ── Styles ───────────────────────────────────────
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="1F3864")
    section_fill = {
        "hazard":   PatternFill("solid", fgColor="D6E4F0"),
        "sec":      PatternFill("solid", fgColor="D5E8D4"),
        "asil":     PatternFill("solid", fgColor="FFE6CC"),
        "req":      PatternFill("solid", fgColor="E1D5E7"),
        "status":   PatternFill("solid", fgColor="F8CECC"),
    }
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap = Alignment(vertical="top", wrap_text=True)
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # ── Headers ──────────────────────────────────────
    headers = [
        ("ID",                    8,  "hazard"),
        ("Hazard",               20,  "hazard"),
        ("Hazardous Event",      30,  "hazard"),
        ("Operational Situation",25,  "hazard"),
        ("Affected Function",    20,  "hazard"),
        ("Source Evidence",      25,  "hazard"),
        ("Severity (S)",         12,  "sec"),
        ("Exposure (E)",         12,  "sec"),
        ("Controllability (C)",  14,  "sec"),
        ("ASIL",                 10,  "asil"),
        ("Safety Goal",          30,  "req"),
        ("FSR",                  35,  "req"),
        ("TSR",                  35,  "req"),
        ("FSR Status",           14,  "status"),
        ("TSR Status",           14,  "status"),
    ]

    for col_idx, (header, width, section) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin
        ws.column_dimensions[cell.column_letter].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    # ── Data rows ─────────────────────────────────────
    asil_colors = {"QM": "C0C0C0", "A": "FFFF99", "B": "FFD700", "C": "FF8C00", "D": "FF0000"}
    status_colors = {"Approved": "00B050", "Rejected": "FF0000", "Edited": "FF8C00"}

    for row_idx, row in enumerate(req.rows, start=2):
        values = [
            row.id, row.hazard, row.hazardous_event, row.operational_situation,
            row.affected_function, row.source_evidence,
            row.S, row.E, row.C, row.asil,
            row.safety_goal, row.fsr, row.tsr,
            row.fsr_status, row.tsr_status,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin
            cell.alignment = wrap

            # Color ASIL column
            if col_idx == 10:
                asil_key = row.asil.replace("ASIL ", "")
                color = asil_colors.get(asil_key, "FFFFFF")
                cell.fill = PatternFill("solid", fgColor=color)
                cell.alignment = center
                cell.font = Font(bold=True)

            # Color status columns
            if col_idx in (14, 15):
                color = status_colors.get(value, "FFFFFF")
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = center

        ws.row_dimensions[row_idx].height = 60

    # Save to buffer and stream back
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=HARA_Traceability_Matrix.xlsx"},
    )


@app.on_event("shutdown")
def cleanup_sessions():
    """Clean up FAISS index directories on shutdown."""
    for path in sessions.values():
        if os.path.exists(path):
            shutil.rmtree(path, ignore_errors=True)